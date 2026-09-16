
import os
import re

from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


# ============================================================
# HELPERS
# ============================================================

def safe_text(value):

    if value is None:
        return ""

    return str(value).strip()


def to_list(value):

    if value is None:
        return []

    if isinstance(
        value,
        (list, tuple, set)
    ):
        return [
            str(item).strip()
            for item in value
            if str(item).strip()
        ]

    text = str(value).strip()

    if not text:
        return []

    return [
        item.strip()
        for item in re.split(
            r"[\n,;]+",
            text
        )
        if item.strip()
    ]


def safe_filename(name):

    name = safe_text(name)

    if not name:
        name = "Generated_CV"

    name = re.sub(
        r"[^A-Za-z0-9_\- ]+",
        "",
        name
    )

    name = name.replace(
        " ",
        "_"
    )

    return name


# ============================================================
# BUILD SCANNER CANDIDATE
# ============================================================

def build_candidate_from_profile(
    profile
):

    full_name = safe_text(
        profile.get(
            "full_name"
        )
    )

    skills = to_list(
        profile.get(
            "skills"
        )
    )

    summary = safe_text(
        profile.get(
            "professional_summary"
        )
    )


    # --------------------------------------------------------
    # Education
    # --------------------------------------------------------

    education_items = profile.get(
        "education",
        []
    )

    if not isinstance(
        education_items,
        list
    ):
        education_items = []


    education_texts = []


    for item in education_items:

        if not isinstance(
            item,
            dict
        ):
            continue

        degree = safe_text(
            item.get(
                "degree"
            )
        )

        field = safe_text(
            item.get(
                "field"
            )
        )

        institution = safe_text(
            item.get(
                "institution"
            )
        )


        parts = [
            part
            for part in [
                degree,
                field,
                institution
            ]
            if part
        ]


        if parts:

            education_texts.append(
                " - ".join(
                    parts
                )
            )


    education_string = (
        " | ".join(
            education_texts
        )
    )


    # --------------------------------------------------------
    # Experience
    # --------------------------------------------------------

    experience_items = profile.get(
        "experience",
        []
    )

    if not isinstance(
        experience_items,
        list
    ):
        experience_items = []


    positions = []

    experience_texts = []


    for item in experience_items:

        if not isinstance(
            item,
            dict
        ):
            continue


        position = safe_text(
            item.get(
                "position"
            )
        )

        company = safe_text(
            item.get(
                "company"
            )
        )

        description = safe_text(
            item.get(
                "description"
            )
        )


        if position:

            positions.append(
                position
            )


        experience_parts = [

            part

            for part in [
                position,
                company,
                description
            ]

            if part
        ]


        if experience_parts:

            experience_texts.append(
                " - ".join(
                    experience_parts
                )
            )


    # --------------------------------------------------------
    # Experience Years
    # --------------------------------------------------------

    try:

        experience_years = float(
            profile.get(
                "experience_years",
                0
            )
        )

    except:

        experience_years = 0.0


    # --------------------------------------------------------
    # Projects
    # --------------------------------------------------------

    project_items = profile.get(
        "projects",
        []
    )

    project_texts = []


    if isinstance(
        project_items,
        list
    ):

        for project in project_items:

            if isinstance(
                project,
                dict
            ):

                name = safe_text(
                    project.get(
                        "name"
                    )
                )

                description = safe_text(
                    project.get(
                        "description"
                    )
                )


                text = " - ".join(
                    part
                    for part in [
                        name,
                        description
                    ]
                    if part
                )


                if text:
                    project_texts.append(
                        text
                    )


    # --------------------------------------------------------
    # Certifications
    # --------------------------------------------------------

    certifications = to_list(
        profile.get(
            "certifications"
        )
    )


    # --------------------------------------------------------
    # Languages
    # --------------------------------------------------------

    languages = to_list(
        profile.get(
            "languages"
        )
    )


    # --------------------------------------------------------
    # Candidate Full Text
    # --------------------------------------------------------

    candidate_text = f"""
Name:
{full_name}

Professional Summary:
{summary}

Skills:
{", ".join(skills)}

Education:
{" | ".join(education_texts)}

Experience:
{" | ".join(experience_texts)}

Projects:
{" | ".join(project_texts)}

Certifications:
{", ".join(certifications)}

Languages:
{", ".join(languages)}
""".strip()


    candidate_id = (
        full_name
        if full_name
        else
        "GENERATED_CV"
    )


    return {

        "candidate_id":
            candidate_id,

        "skills":
            skills,

        "education":
            education_string,

        "experience_years":
            experience_years,

        "positions":
            positions,

        "summary":
            summary,

        "candidate_text":
            candidate_text
    }


# ============================================================
# DOCX HELPERS
# ============================================================

def add_section_title(
    document,
    title
):

    paragraph = document.add_paragraph()

    run = paragraph.add_run(
        title
    )

    run.bold = True
    run.font.size = Pt(
        13
    )


def add_bullet(
    document,
    text
):

    text = safe_text(
        text
    )

    if not text:
        return

    paragraph = document.add_paragraph(
        style="List Bullet"
    )

    paragraph.add_run(
        text
    )


# ============================================================
# GENERATE PROFESSIONAL CV
# ============================================================

def generate_cv_docx(
    profile,
    output_directory,
    filename=None
):

    output_directory = Path(
        output_directory
    )

    output_directory.mkdir(
        parents=True,
        exist_ok=True
    )


    full_name = safe_text(
        profile.get(
            "full_name"
        )
    )


    if filename is None:

        filename = (
            safe_filename(
                full_name
            )
            +
            "_CV.docx"
        )


    if not filename.lower().endswith(
        ".docx"
    ):

        filename += ".docx"


    output_path = (
        output_directory
        /
        filename
    )


    document = Document()


    # ========================================================
    # PAGE SETTINGS
    # ========================================================

    section = document.sections[0]

    section.top_margin = Inches(
        0.6
    )

    section.bottom_margin = Inches(
        0.6
    )

    section.left_margin = Inches(
        0.7
    )

    section.right_margin = Inches(
        0.7
    )


    # ========================================================
    # NAME
    # ========================================================

    name_paragraph = (
        document.add_paragraph()
    )

    name_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )


    name_run = (
        name_paragraph.add_run(
            full_name
            or
            "Candidate Name"
        )
    )

    name_run.bold = True

    name_run.font.size = Pt(
        20
    )


    # ========================================================
    # CONTACT INFO
    # ========================================================

    contact_items = [

        safe_text(
            profile.get(
                "email"
            )
        ),

        safe_text(
            profile.get(
                "phone"
            )
        ),

        safe_text(
            profile.get(
                "location"
            )
        ),

        safe_text(
            profile.get(
                "linkedin"
            )
        ),

        safe_text(
            profile.get(
                "github"
            )
        )
    ]


    contact_items = [
        item
        for item in contact_items
        if item
    ]


    if contact_items:

        contact_paragraph = (
            document.add_paragraph()
        )

        contact_paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        contact_paragraph.add_run(
            " | ".join(
                contact_items
            )
        )


    # ========================================================
    # SUMMARY
    # ========================================================

    summary = safe_text(
        profile.get(
            "professional_summary"
        )
    )


    if summary:

        add_section_title(
            document,
            "Professional Summary"
        )

        document.add_paragraph(
            summary
        )


    # ========================================================
    # SKILLS
    # ========================================================

    skills = to_list(
        profile.get(
            "skills"
        )
    )


    if skills:

        add_section_title(
            document,
            "Skills"
        )

        document.add_paragraph(
            " • ".join(
                skills
            )
        )


    # ========================================================
    # EXPERIENCE
    # ========================================================

    experience_items = profile.get(
        "experience",
        []
    )


    if (
        isinstance(
            experience_items,
            list
        )
        and
        experience_items
    ):

        add_section_title(
            document,
            "Professional Experience"
        )


        for item in experience_items:

            if not isinstance(
                item,
                dict
            ):
                continue


            position = safe_text(
                item.get(
                    "position"
                )
            )

            company = safe_text(
                item.get(
                    "company"
                )
            )

            start_date = safe_text(
                item.get(
                    "start_date"
                )
            )

            end_date = safe_text(
                item.get(
                    "end_date"
                )
            )

            description = safe_text(
                item.get(
                    "description"
                )
            )


            heading = " | ".join(
                part
                for part in [
                    position,
                    company
                ]
                if part
            )


            if heading:

                paragraph = (
                    document.add_paragraph()
                )

                run = paragraph.add_run(
                    heading
                )

                run.bold = True


            dates = " - ".join(
                part
                for part in [
                    start_date,
                    end_date
                ]
                if part
            )


            if dates:

                document.add_paragraph(
                    dates
                )


            if description:

                description_lines = [

                    line.strip()

                    for line
                    in description.split(
                        "\n"
                    )

                    if line.strip()
                ]


                for line in description_lines:

                    add_bullet(
                        document,
                        line
                    )


    # ========================================================
    # EDUCATION
    # ========================================================

    education_items = profile.get(
        "education",
        []
    )


    if (
        isinstance(
            education_items,
            list
        )
        and
        education_items
    ):

        add_section_title(
            document,
            "Education"
        )


        for item in education_items:

            if not isinstance(
                item,
                dict
            ):
                continue


            degree = safe_text(
                item.get(
                    "degree"
                )
            )

            field = safe_text(
                item.get(
                    "field"
                )
            )

            institution = safe_text(
                item.get(
                    "institution"
                )
            )

            year = safe_text(
                item.get(
                    "year"
                )
            )


            line = " | ".join(
                part
                for part in [
                    degree,
                    field,
                    institution,
                    year
                ]
                if part
            )


            if line:

                add_bullet(
                    document,
                    line
                )


    # ========================================================
    # PROJECTS
    # ========================================================

    project_items = profile.get(
        "projects",
        []
    )


    if (
        isinstance(
            project_items,
            list
        )
        and
        project_items
    ):

        add_section_title(
            document,
            "Projects"
        )


        for project in project_items:

            if not isinstance(
                project,
                dict
            ):
                continue


            project_name = safe_text(
                project.get(
                    "name"
                )
            )

            project_description = (
                safe_text(
                    project.get(
                        "description"
                    )
                )
            )


            if project_name:

                paragraph = (
                    document.add_paragraph()
                )

                run = paragraph.add_run(
                    project_name
                )

                run.bold = True


            if project_description:

                add_bullet(
                    document,
                    project_description
                )


    # ========================================================
    # CERTIFICATIONS
    # ========================================================

    certifications = to_list(
        profile.get(
            "certifications"
        )
    )


    if certifications:

        add_section_title(
            document,
            "Certifications"
        )


        for certificate in certifications:

            add_bullet(
                document,
                certificate
            )


    # ========================================================
    # LANGUAGES
    # ========================================================

    languages = to_list(
        profile.get(
            "languages"
        )
    )


    if languages:

        add_section_title(
            document,
            "Languages"
        )

        document.add_paragraph(
            " • ".join(
                languages
            )
        )


    document.save(
        output_path
    )


    return str(
        output_path
    )
